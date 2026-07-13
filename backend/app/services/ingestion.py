import os
import shutil
import logging
from typing import List, Dict, Any
from PIL import Image
import fitz  # PyMuPDF
import docx  # python-docx
from pptx import Presentation  # python-pptx
import pytesseract

logger = logging.getLogger(__name__)

# Fallback: check standard Tesseract installation paths on Windows if not in PATH
TESSERACT_DEFAULT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
if shutil.which("tesseract") is None:
    if os.path.exists(TESSERACT_DEFAULT_PATH):
        pytesseract.pytesseract.tesseract_cmd = TESSERACT_DEFAULT_PATH
        logger.info(
            f"Tesseract binary not found in system PATH. Pointed pytesseract to: {TESSERACT_DEFAULT_PATH}"
        )
    else:
        logger.warning(
            "Tesseract binary not found in system PATH and not found at default location. "
            "OCR fallback will be unavailable unless Tesseract is installed."
        )


class IngestionService:
    def extract_text(self, file_path: str, file_extension: str) -> List[Dict[str, Any]]:
        """
        Dispatches to the correct parser based on file extension.
        Returns a list of dictionaries containing page number and text:
        [{"page_number": 1, "text": "..."}]
        """
        ext = file_extension.lower().strip(".")

        if ext == "pdf":
            return self._extract_pdf(file_path)
        elif ext in ["doc", "docx"]:
            return self._extract_docx(file_path)
        elif ext in ["ppt", "pptx"]:
            return self._extract_pptx(file_path)
        elif ext == "txt":
            return self._extract_txt(file_path)
        else:
            raise ValueError(f"Unsupported file extension: .{ext}")

    def _extract_pdf(self, file_path: str) -> List[Dict[str, Any]]:
        """
        Extracts text from PDF using PyMuPDF. Falls back to Tesseract OCR for scanned pages.
        """
        pages_data = []
        doc = fitz.open(file_path)

        for page_index in range(len(doc)):
            page_num = page_index + 1
            page = doc.load_page(page_index)
            text = page.get_text().strip()

            # If the page has very little or no text, it's likely scanned. Attempt OCR fallback.
            if len(text) < 50:
                logger.info(
                    f"Page {page_num} text density low ({len(text)} chars). Attempting Tesseract OCR fallback."
                )
                try:
                    # Render page to a high quality image (DPI=150 is a good speed/accuracy balance)
                    pix = page.get_pixmap(dpi=150)

                    # Convert PyMuPDF pixmap to PIL Image
                    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

                    # Run Tesseract OCR on the image
                    ocr_text = pytesseract.image_to_string(img).strip()
                    if ocr_text:
                        text = f"[OCR Extracted]\n{ocr_text}"
                        logger.info(f"Page {page_num} OCR successful.")
                    else:
                        text = "[Scanned Page - No text found via OCR]"
                        logger.warning(f"Page {page_num} OCR returned empty string.")

                except pytesseract.TesseractNotFoundError:
                    text = "[Scanned Page - Tesseract OCR not installed on server]"
                    logger.error(
                        f"Tesseract OCR is not installed or not in PATH. Skipping OCR for page {page_num}."
                    )
                except Exception as e:
                    text = f"[Scanned Page - OCR Error: {str(e)}]"
                    logger.error(
                        f"Error during OCR extraction on page {page_num}: {str(e)}",
                        exc_info=True,
                    )

            pages_data.append({"page_number": page_num, "text": text})

        doc.close()
        return pages_data

    def _extract_docx(self, file_path: str) -> List[Dict[str, Any]]:
        """
        Extracts text from Word documents using python-docx.
        Since flow documents do not have physical page numbers, all text goes into page 1.
        """
        doc = docx.Document(file_path)
        full_text = []

        # Extract text from paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)

        # Extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                row_text = [
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                ]
                if row_text:
                    full_text.append(" | ".join(row_text))

        return [{"page_number": 1, "text": "\n".join(full_text)}]

    def _extract_pptx(self, file_path: str) -> List[Dict[str, Any]]:
        """
        Extracts text from PowerPoint presentations using python-pptx.
        Each slide maps to a "page".
        """
        prs = Presentation(file_path)
        pages_data = []

        for index, slide in enumerate(prs.slides):
            page_num = index + 1
            slide_text = []

            for shape in slide.shapes:
                if hasattr(shape, "text_frame") and shape.text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            slide_text.append(paragraph.text)

            pages_data.append(
                {"page_number": page_num, "text": "\n".join(slide_text).strip()}
            )

        return pages_data

    def _extract_txt(self, file_path: str) -> List[Dict[str, Any]]:
        """
        Reads a standard plain text file. Entire text goes into page 1.
        """
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()

        return [{"page_number": 1, "text": text.strip()}]


# Instantiate singleton service instance
ingestion_service = IngestionService()
